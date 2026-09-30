# -*- coding: utf-8 -*-
"""Reference 16:9 deck in the AGM system (Seafood and Marine editions).

Five layouts, each shown once: cover, section divider, content, two-column,
and closing. Sample content only -- swap SLIDES and the layouts hold.
"""
import os
import tempfile

from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from brands import BRANDS, DEFAULT
from palette import F400, F500, F600, F800, SLIDE

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
BR = BRANDS[DEFAULT]          # active edition (see use())
P = BR.P
_BANDS = {}


def use(edition):
    """Switch the active edition: 'seafood' or 'marine'."""
    global BR, P
    BR = BRANDS[edition]
    P = BR.P


def deck_band():
    """The full-bleed device band for the active edition (generated once)."""
    if BR.key not in _BANDS:
        out = os.path.join(tempfile.mkdtemp(), "deck_band_%s.png" % BR.key)
        _BANDS[BR.key] = BR.deck_band(out)
    return _BANDS[BR.key]
SW, SH = SLIDE["w"], SLIDE["h"]
ML, MR = SLIDE["ml"], SLIDE["mr"]
CW = SW - ML - MR                      # 302.667mm content measure


def rgb(h):
    return RGBColor.from_string(h)


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def box(slide, x, y, w, h, fill=None, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Mm(x), Mm(y), Mm(w), Mm(h))
    if radius:
        # adj is the corner radius as a fraction of the shorter side
        shp.adjustments[0] = min(radius, min(w, h) / 2.0) / min(w, h)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def line(tf, first=False, space_before=0, space_after=0, align=PP_ALIGN.LEFT,
         spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if spacing:
        p.line_spacing = spacing
    return p


def txt(p, text, font=F400, size=14, color="000000", tr=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = font
    r.font.color.rgb = rgb(color)
    if tr is not None:                  # letter spacing, in points
        r.font._rPr.set("spc", str(int(round(tr * 100))))
    return r


def rule(slide, x, y, w, h, color):
    return box(slide, x, y, w, h, color)


# ------------------------------------------------------------------ slides --
def cover(prs, title, subtitle, date):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, P["PANEL"])
    s.shapes.add_picture(deck_band(),
                         Mm(0), Mm(SH - 52), Mm(SW), Mm(52))
    s.shapes.add_picture(BR.emblem_white,
                         Mm(ML), Mm(20), height=Mm(18))

    tf = textbox(s, ML, 48, CW * 0.74, 60)
    p = line(tf, first=True)
    txt(p, BR.wordmark, F800, 11, P["PALE"], 1.1)
    p = line(tf, space_before=10, spacing=0.92)
    txt(p, title, F800, 40, P["PAPER"], -0.8)
    p = line(tf, space_before=8)
    txt(p, subtitle, F500, 15, P["PALE"])

    tf = textbox(s, SW - MR - 70, 24, 70, 8)
    p = line(tf, first=True, align=PP_ALIGN.RIGHT)
    txt(p, date, F600, 9.5, P["PALE"], 1.4)
    return s


def divider(prs, number, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, P["PANEL"])
    rule(s, 0, 0, SW, 1.6, P["ACCENT"])

    tf = textbox(s, ML, SH / 2 - 26, CW, 52, MSO_ANCHOR.MIDDLE)
    p = line(tf, first=True)
    txt(p, number, F800, 54, P["META"], -1.0)
    p = line(tf, space_before=4)
    txt(p, title, F800, 34, P["PAPER"], -0.7)
    p = line(tf, space_before=6)
    txt(p, subtitle, F500, 14, P["PALE"])
    return s


def content(prs, eyebrow, title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, P["PAPER"])

    tf = textbox(s, ML, 20, CW, 26)
    p = line(tf, first=True)
    txt(p, eyebrow, F800, 9.5, P["META"], 1.6)
    p = line(tf, space_before=5)
    txt(p, title, F800, 28, P["INK"], -0.6)
    rule(s, ML, 54, 34, 1.2, P["ACCENT"])

    tf = textbox(s, ML, 64, CW * 0.82, 90)
    for i, b in enumerate(bullets):
        p = line(tf, first=(i == 0), space_before=0 if i == 0 else 11,
                 spacing=1.35)
        txt(p, "—   ", F800, 13, P["RULE"])
        txt(p, b, F400, 13, P["BODY"])
    footer(s, prs)
    return s


def two_col(prs, eyebrow, title, cards):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, P["PAPER"])

    tf = textbox(s, ML, 20, CW, 26)
    p = line(tf, first=True)
    txt(p, eyebrow, F800, 9.5, P["META"], 1.6)
    p = line(tf, space_before=5)
    txt(p, title, F800, 28, P["INK"], -0.6)
    rule(s, ML, 54, 34, 1.2, P["ACCENT"])

    gap, n = 8.0, len(cards)
    cw = (CW - gap * (n - 1)) / n
    top, ch = 70.0, 56.0
    for i, (head, body) in enumerate(cards):
        x = ML + i * (cw + gap)
        box(s, x, top, cw, ch, P["TINT"], radius=3.0)
        rule(s, x + 9, top + 8, 16, 1.2, P["ACCENT"])   # inside the card, not on its rim
        tf = textbox(s, x + 9, top + 16, cw - 18, ch - 24)
        p = line(tf, first=True)
        txt(p, head, F800, 15, P["INK"], -0.3)
        p = line(tf, space_before=7, spacing=1.3)
        txt(p, body, F400, 11, P["BODY"])
    footer(s, prs)
    return s


def closing(prs, name, role, mail, tel, prefix=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, P["PANEL"])
    s.shapes.add_picture(deck_band(),
                         Mm(0), Mm(SH - 52), Mm(SW), Mm(52))
    s.shapes.add_picture(BR.emblem_white,
                         Mm(ML), Mm(20), height=Mm(18))

    tf = textbox(s, ML, 52, CW * 0.7, 50)
    p = line(tf, first=True)
    txt(p, "YOUR CONTACT AT " + BR.name.upper(), F800, 9, P["PALE"], 1.8)
    p = line(tf, space_before=9)
    if prefix:
        txt(p, prefix, F500, 18, P["PALE"])
    txt(p, name, F800, 26, P["PAPER"], -0.5)
    p = line(tf, space_before=7)
    txt(p, role, F600, 12.5, P["TINT"])
    txt(p, "   ·   ", F400, 12.5, P["PALE"])
    txt(p, BR.name, F500, 12.5, P["PALE"])
    p = line(tf, space_before=4)
    txt(p, mail, F600, 12.5, P["TINT"])
    txt(p, "   ·   ", F400, 12.5, P["PALE"])
    txt(p, tel, F400, 12.5, P["PALE"])
    return s


def footer(s, prs):
    rule(s, ML, SH - 14, SW - ML - MR, 0.35, P["RULE"])
    tf = textbox(s, ML, SH - 11, CW * 0.7, 6)
    p = line(tf, first=True)
    txt(p, BR.footer_tagline, F500, 8, P["MUTED"])
    tf = textbox(s, SW - MR - 40, SH - 11, 40, 6)
    p = line(tf, first=True, align=PP_ALIGN.RIGHT)
    txt(p, BR.name, F600, 8, P["MUTED"])


def sample(prs):
    """Reference deck: every layout once. Write your own compose(prs) that
    calls cover / divider / content / two_col / closing and pass it to build()."""
    cover(prs, "Deck Title Goes Here",
          "Subtitle line — what this deck is for and who it is for",
          "24 SEPTEMBER 2026")
    divider(prs, "01", "Section Title", "What this section covers")
    content(prs, "SECTION EYEBROW", "Headline that states the point", [
        "First supporting point. One idea per line, written as a statement "
        "rather than a label.",
        "Second supporting point, long enough to show how a line wraps "
        "against the measure and keeps its leading.",
        "Third supporting point.",
    ])
    two_col(prs, "SECTION EYEBROW", "Three things side by side", [
        ("Regulatory", "Permits, licensing and the approvals pathway, "
                       "mapped against the responsible ministry."),
        ("Commercial", "Cost structure, who carries which cost, and where "
                       "the commercial risk actually sits."),
        ("Delivery", "Sequence, milestones and the decisions that have to "
                     "be made before each one."),
    ])
    closing(prs, "Name Surname", "Role title",
            BR.contact["email"], BR.contact["phone"])


def build(out, compose=sample, edition=DEFAULT):
    use(edition)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(SW), Mm(SH)
    compose(prs)
    prs.save(out)
    return out


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Build the AGM reference deck (or import build() with your own compose(prs)).")
    ap.add_argument("--out", default="AGM_Deck.pptx")
    ap.add_argument("--edition", default=DEFAULT, choices=sorted(BRANDS))
    a = ap.parse_args()
    print(build(a.out, edition=a.edition))
