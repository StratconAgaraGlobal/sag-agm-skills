#!/usr/bin/env python3
"""Build every reference output for both brands and check the files exist.

    python tests/smoke_test.py

Needs python-docx, python-pptx and Pillow. Fonts are not required to build
(they are embedded / only matter for rendering).
"""
import json
import os
import re
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SK = os.path.join(ROOT, "skills")

JOBS = [
    ("sag-brand", ["build_doc.py", "--out", "{t}/sag_doc.docx"]),
    ("sag-brand", ["build_mom.py", "--out", "{t}/sag_mom.docx"]),
    ("sag-brand", ["build_deck.py", "--out", "{t}/sag_deck.pptx"]),
    ("sag-brand", ["build_example_html.py", "--out", "{t}/sag.html"]),
]
for ed in ("seafood", "marine"):
    JOBS += [
        ("agm-brand", ["build_doc.py", "--edition", ed, "--out", "{t}/agm_doc_%s.docx" % ed]),
        ("agm-brand", ["build_mom.py", "--brand", ed, "--out", "{t}/agm_mom_%s.docx" % ed]),
        ("agm-brand", ["build_deck.py", "--edition", ed, "--out", "{t}/agm_deck_%s.pptx" % ed]),
    ]
JOBS.append(("agm-brand", ["build_example_html.py", "--out", "{t}/agm.html"]))

# The SAG deck's stat-card and case-study layouts, beyond the reference deck:
# every row count, photos (cropped to fill), a logo and the error paths.
SAG_LAYOUTS = r"""
import sys
from PIL import Image
from pptx import Presentation
import build_deck as b

t = sys.argv[1]
photo, tall, logo = t + "/photo.jpg", t + "/tall.jpg", t + "/logo.png"
Image.new("RGB", (1600, 1000), (90, 120, 130)).save(photo)
Image.new("RGB", (800, 1200), (140, 110, 90)).save(tall)
Image.new("RGB", (600, 150), (27, 48, 56)).save(logo)

def compose(prs):
    for n in range(1, 9):
        b.stat_cards(prs, "E", "%d stats" % n,
                     [("IDR 1.5T+", "label %d" % i) for i in range(n)])
    body = "Sentence long enough to wrap across the column. " * 3
    for photos in ((), (photo,), (photo, tall)):
        for n in (0, 1, 2, 4):
            if not photos and not n:
                continue
            b.case_study(prs, "CASE STUDY", "Headline", "Client",
                         "Place  ·  Sector  ·  Years", body, body, body,
                         stats=[("700K+ m\u00b2", "land")] * n,
                         photos=photos, logo=logo if n % 2 else None)

out = b.build(t + "/sag_layouts.pptx", compose)
prs = Presentation(out)
assert len(prs.slides) == 8 + 11, len(prs.slides)
pics = sum(sh.shape_type == 13 for s in prs.slides for sh in s.shapes)
assert pics == 4 + 2 * 4 + 3, pics       # photos, plus a logo where n is odd
for s in prs.slides:
    for sh in s.shapes:
        assert sh.left >= 0 and sh.top >= 0, sh.name
        assert sh.left + sh.width <= prs.slide_width, sh.name
        assert sh.top + sh.height <= prs.slide_height, sh.name

ref = Presentation(b.build(t + "/sag_ref.pptx"))
assert len(ref.slides) == 7, len(ref.slides)

for bad in (lambda p: b.stat_cards(p, "E", "T", []),
            lambda p: b.stat_cards(p, "E", "T", [("1", "x")] * 9),
            lambda p: b.case_study(p, "E", "T", "C", "M", "a", "b", "c",
                                   stats=[("1", "x")] * 5)):
    try:
        b.build(t + "/bad.pptx", bad)
    except ValueError:
        continue
    raise AssertionError("expected ValueError")
print("ok")
"""


def check_repo():
    """The plugin manifest lists every skill, and the app files reference only what exists."""
    fails = 0
    m = json.load(open(os.path.join(ROOT, ".claude-plugin", "marketplace.json"), encoding="utf-8"))
    listed = {os.path.normpath(p) for e in m["plugins"] for p in e.get("skills", [])}
    on_disk = {os.path.normpath("./skills/" + d) for d in os.listdir(SK)
               if os.path.isfile(os.path.join(SK, d, "SKILL.md"))}
    ok = listed == on_disk and all("version" not in e for e in m["plugins"])
    print("%-4s %s" % ("ok" if ok else "FAIL", "marketplace.json lists every skill, no pinned version"))
    if not ok:
        fails += 1
        print("  missing:", sorted(on_disk - listed), " extra:", sorted(listed - on_disk))
    ex = os.path.join(SK, "sag-brand", "examples", "sag-app-example.html")
    refs = re.findall(r"""(?:href|src)="(\.\./[^"]+)"|url\('(\.\./[^']+)'\)""", open(ex, encoding="utf-8").read())
    missing = [p for pair in refs for p in pair
               if p and not os.path.exists(os.path.normpath(os.path.join(os.path.dirname(ex), p)))]
    print("%-4s %s" % ("FAIL" if missing else "ok", "sag-app-example.html local references"))
    if missing:
        fails += 1
        print("  missing:", missing)
    return fails


def main():
    fails = check_repo()
    with tempfile.TemporaryDirectory() as t:
        for skill, args in JOBS:
            args = [a.replace("{t}", t) for a in args]
            cwd = os.path.join(SK, skill, "scripts")
            r = subprocess.run([sys.executable] + args, cwd=cwd, capture_output=True, text=True)
            out = args[args.index("--out") + 1]
            ok = r.returncode == 0 and os.path.getsize(out) > 5000
            print("%-4s %s %s" % ("ok" if ok else "FAIL", skill, " ".join(args[:-2])))
            if not ok:
                fails += 1
                print(r.stdout[-500:], r.stderr[-1500:])
        r = subprocess.run([sys.executable, "-c", SAG_LAYOUTS, t],
                           cwd=os.path.join(SK, "sag-brand", "scripts"),
                           capture_output=True, text=True)
        ok = r.returncode == 0 and r.stdout.strip() == "ok"
        print("%-4s %s %s" % ("ok" if ok else "FAIL", "sag-brand", "deck layouts"))
        if not ok:
            fails += 1
            print(r.stdout[-500:], r.stderr[-1500:])
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
