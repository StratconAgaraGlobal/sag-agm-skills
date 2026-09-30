# -*- coding: utf-8 -*-
"""BrandDoc - the SAG component system for Word, brand-agnostic.

One builder for SAG and for AGM (both editions): the geometry, type and
component sizes are identical, only the `brand` object differs (palette,
wordmark, logo files, banner images, footer tagline). Everything is laid out
to DESIGN_SPEC.md: A4, 20/20/14/16 mm margins, 170 mm measure, 2.5 mm radius.

    from brands import SAG            # or SEAFOOD / MARINE in agm-brand
    from components import BrandDoc

    d = BrandDoc(SAG, kind="Minutes of Meeting", title="Short title",
                 date="29 SEPTEMBER 2026")
    d.masthead()
    d.title_block("MINUTES OF MEETING · REF · v0.9 DRAFT", "Title",
                  "Subtitle lead ", "counterparty")
    d.intro(["Paragraph with **bold** runs."])
    d.tiles([("0%", "TRADING FEE", "Free to use"), ...])
    d.callout("HEADLINE FOR MANAGEMENT", ["**Opportunity.** ..."])
    d.band("Meeting details", "Who, when and why")
    d.kv([("Date", "Tuesday 29 September 2026"), ...])
    d.data_table(["Organisation", "Name", "Role"], rows, [55, 45, 70])
    d.save("out.docx")

`**text**` inside any string renders as SemiBold in the ink colour.
Shapes (tiles, callouts, bands, panels) are fixed-height DrawingML shapes and
do not grow, so callout() estimates its height from the text.
"""
import math
import os
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Mm, Pt, RGBColor

from palette import F400, F500, F600, F800, PAGE, W, GUTTER, RADIUS
from sagdoc import (mm, track, el, run, fmt, para, cell_p, shade, cborders,
                    cmargin, valign, rowheight, cantsplit, keeprow, table,
                    addrow, spacer, drop_empty_first, shape, emit, xrun,
                    xpara, embed_fonts, picture, repeat_header, pborder)
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = os.path.join(os.path.dirname(HERE), "assets", "fonts")
LEFT, RIGHT, CENTER = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT,
                       WD_ALIGN_PARAGRAPH.CENTER)
BODY_W = W - GUTTER
ROW_W = W - 0.4            # shape rows total 169.6 mm so they never wrap
LINE_MM = 5.7              # one 9.2 pt line at 1.4 leading, as Word/LibreOffice render it


# ------------------------------------------------------------ rich text ----
def _segments(text):
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]"""
    out, bold = [], False
    for i, part in enumerate(text.split("**")):
        if part:
            out.append((part, bold))
        bold = not bold
    return out


def rich(p, text, font=F400, size=10, color="000000", bold_font=F600,
         bold_color=None, tr=0):
    for seg, b in _segments(text):
        run(p, seg, font=bold_font if b else font, size=size,
            color=(bold_color or color) if b else color, tr=tr)


def xrich(text, font, size, color, bold_font=F600, bold_color=None, tr=0):
    return [xrun(seg, bold_font if b else font, size,
                 (bold_color or color) if b else color, tr)
            for seg, b in _segments(text)]


def _plain(text):
    return text.replace("**", "")


def est_lines(text, width_mm, size_pt, factor=0.49):
    """Conservative line count for text set in Plus Jakarta Sans."""
    char_mm = size_pt * 0.3528 * factor
    per_line = max(10, int(width_mm / char_mm))
    words, lines, cur = _plain(text).split(), 1, 0
    for w_ in words:
        if cur + len(w_) + (1 if cur else 0) > per_line:
            lines, cur = lines + 1, len(w_)
        else:
            cur += len(w_) + (1 if cur else 0)
    return lines


MARK_ROLE = {"HIGH": "INK", "MEDIUM": "LBL", "LOW": "MUTED",
             "OPEN": "LBL", "PENDING": "LBL", "DONE": "MUTED",
             "TBC": "LBL", "TO DO": "LBL", "IN PROGRESS": "LBL",
             "CLOSED": "MUTED", "YES": "INK", "NO": "MUTED"}


class BrandDoc:
    """Builds one A4 document. See the module docstring."""

    def __init__(self, brand, kind, title, date, embed=True):
        self.b = brand
        self.P = brand.P
        self.kind, self.title, self.date = kind, title, date
        self.embed = embed
        self.tmp = tempfile.mkdtemp()
        self.num = 0
        P = self.P
        self.LBL = getattr(brand, "LBL", None) or P["META"]
        self.ON_DARK = P.get("ACCENT_ON_DARK", P["ACCENT"])

        self.doc = doc = Document()
        sec = self.sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(PAGE["w"]), Mm(PAGE["h"])
        sec.left_margin, sec.right_margin = Mm(PAGE["ml"]), Mm(PAGE["mr"])
        sec.top_margin, sec.bottom_margin = Mm(PAGE["mt"]), Mm(PAGE["mb"])
        sec.header_distance = Mm(PAGE["header"])
        sec.footer_distance = Mm(PAGE["footer"])
        doc.element.insert(0, el("w:background", color=P["PAPER"]))
        doc.settings.element.append(el("w:displayBackgroundShape"))

        st = doc.styles["Normal"]
        st.font.name = F400
        st.font.size = Pt(10)
        st.font.color.rgb = RGBColor.from_string(P["INK"])
        st.element.rPr.rFonts.set(qn("w:hAnsi"), F400)
        st.element.rPr.rFonts.set(qn("w:cs"), F400)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        self._running()

    # ------------------------------------------------------- page 1 furniture
    def masthead(self, date=None):
        out = os.path.join(self.tmp, "masthead.png")
        self.b.masthead(out, date or self.date)
        picture(para(self.doc, align=LEFT), out, width_mm=W)

    def title_block(self, eyebrow, title, subtitle_lead="", subtitle_name=""):
        P, d = self.P, self.doc
        run(para(d, before=13), eyebrow, font=F800, size=8, color=self.LBL,
            tr=track(8, 0.20))
        run(para(d, before=4), title, font=F800, size=23, color=P["INK"],
            tr=track(23, -0.022))
        p = para(d, before=3.5)
        run(p, subtitle_lead, font=F500, size=11, color=P["BODY"])
        run(p, subtitle_name, font=F600, size=11, color=P["PANEL"])
        spacer(d, 5)
        t = table(d, [17, 17, 17, W - 51])
        rowheight(t.rows[0], 1.1, "exact")
        for c, k in zip(t.rows[0].cells[:3], ("RULE", "PALE", "ACCENT")):
            shade(c, P[k])
        spacer(d, 5)

    def intro(self, paragraphs):
        for i, txt in enumerate(paragraphs):
            p = para(self.doc, before=0 if i == 0 else 5, line=1.45, ind_r=6)
            rich(p, txt, F400, 10, self.P["BODY"], bold_color=self.P["INK"])

    # ------------------------------------------------------------- panels ---
    def tiles(self, items, dark=True, gap=3.0):
        """Row of headline figures: (value, label, sub). Total 169.6 mm."""
        P, n = self.P, len(items)
        w = (ROW_W - gap * (n - 1)) / n
        fill = P["PANEL"] if dark else P["TINT"]
        vcol = self.ON_DARK if dark else P["INK"]
        lcol = P["PAPER"] if dark else P["META"]
        scol = P["PALE"] if dark else P["BODY"]
        parts = []
        for i, (val, label, sub) in enumerate(items):
            body = xpara([xrun(val, F800, 17, vcol, track(17, -0.02))],
                         after=2, jc="center")
            body += xpara([xrun(label.upper(), F800, 6.4, lcol, track(6.4, 0.2))],
                          after=1.5, jc="center")
            if sub:
                body += xpara([xrun(sub, F400, 6.8, scol)], jc="center",
                              line=1.15)
            if i:
                parts.append(shape(gap, 28.0, None, "<w:p/>", 0))
            parts.append(shape(w, 28.0, fill, body, RADIUS,
                               pad=(1.5, 2.6, 1.5, 2.6)))
        emit(self.doc, parts, before=11)

    def callout(self, label, paragraphs, tone="tint", before=9):
        P = self.P
        fill = P["TINT"] if tone == "tint" else P["PANEL"]
        lcol = self.LBL if tone == "tint" else self.ON_DARK
        tcol = P["BODY"] if tone == "tint" else P["PALE"]
        bcol = P["INK"] if tone == "tint" else P["PAPER"]
        inner = W - 12
        body = xpara([xrun(label.upper(), F800, 8, lcol, track(8, 0.18))],
                     after=4)
        h = 10 + 5.0 + 1.5
        for i, txt in enumerate(paragraphs):
            body += xpara(xrich(txt, F400, 9.2, tcol, F600, bcol),
                          before=0 if i == 0 else 3.5, line=1.4)
            h += est_lines(txt, inner, 9.2) * LINE_MM + (0 if i == 0 else 1.2)
        emit(self.doc, [shape(W, round(h, 1), fill, body, RADIUS,
                              pad=(6, 5, 6, 5), anchor="t")], before=before)

    def band(self, title, subtitle="", first=False):
        P = self.P
        spacer(self.doc, 2 if first else 11, keep=True)
        runs = [xrun(title.upper(), F800, 11, P["PAPER"], track(11, 0.09))]
        if subtitle:
            runs += ["<w:r><w:tab/></w:r>", xrun(subtitle, F500, 8.8, P["PALE"])]
        emit(self.doc, [shape(W, 12.7, P["PANEL"],
                              xpara(runs, tab_mm=W - 11.5), RADIUS,
                              pad=(5.5, 2.6, 6, 2.6))], keep=True)
        spacer(self.doc, 5, keep=True)

    def label(self, text, before=8):
        p = para(self.doc, before=before, after=3, keep=True)
        run(p, text.upper(), font=F800, size=8, color=self.LBL,
            tr=track(8, 0.20))

    def note(self, text, before=4):
        p = para(self.doc, before=before, line=1.3)
        rich(p, text, F400, 8, self.P["MUTED"], bold_color=self.P["BODY"])
        for r in p.runs:
            r.italic = True

    def para(self, text, before=5):
        p = para(self.doc, before=before, line=1.4)
        rich(p, text, F400, 9.6, self.P["BODY"], bold_color=self.P["INK"])

    def bullets(self, items, before=1.5):
        for i, txt in enumerate(items):
            p = para(self.doc, before=before if i else 0, line=1.4)
            p.paragraph_format.left_indent = Mm(5)
            p.paragraph_format.first_line_indent = Mm(-5)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(5))
            p.paragraph_format.keep_with_next = i < len(items) - 1
            run(p, "–\t", font=F800, size=10, color=self.LBL)
            rich(p, txt, F400, 9.8, self.P["BODY"], bold_color=self.P["INK"])

    def numbered(self, items, before=1.5):
        for i, txt in enumerate(items):
            p = para(self.doc, before=before if i else 0, line=1.4)
            p.paragraph_format.left_indent = Mm(8)
            p.paragraph_format.first_line_indent = Mm(-8)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(8))
            p.paragraph_format.keep_with_next = i < len(items) - 1
            run(p, "%02d\t" % (i + 1), font=F800, size=9.8, color=self.LBL)
            rich(p, txt, F400, 9.8, self.P["BODY"], bold_color=self.P["INK"])

    # --------------------------------------------------------------- tables -
    def kv(self, rows, key_w=35.0, before=4):
        """Two-column facts table: shaded key cell, plain value cell."""
        P = self.P
        spacer(self.doc, before, keep=True)
        cols = [key_w, W - key_w]
        t = table(self.doc, cols, lmar=3.5)
        for i, (k, v) in enumerate(rows):
            r = t.rows[0] if i == 0 else addrow(t, cols)
            rowheight(r, 7.6)
            cantsplit(r)
            for c in r.cells:
                valign(c, "center")
                cmargin(c, l=(3.5 if c is r.cells[0] else 4), t=1.8, b=1.8)
            shade(r.cells[0], P["TINT"])
            cborders(r.cells[0], bottom=(8, P["PAPER"]))
            run(cell_p(r.cells[0], align=LEFT), k, font=F800, size=8.2,
                color=P["INK"])
            rich(cell_p(r.cells[1], align=LEFT, line=1.25), v, F400, 9.2,
                 P["BODY"], bold_color=P["INK"])

    def data_table(self, headers, rows, widths, align=None, marks=(),
                   first_bold=True, before=4):
        """Header row + hairline rows. `marks` = column indexes that hold
        status marks (OPEN / HIGH / TBC ...), set as small tracked caps."""
        P = self.P
        assert abs(sum(widths) - W) < 0.6, "column widths must total %.0f mm" % W
        align = align or ["l"] * len(headers)
        A = {"l": LEFT, "c": CENTER, "r": RIGHT}
        spacer(self.doc, before, keep=True)
        t = table(self.doc, widths)
        r0 = t.rows[0]
        repeat_header(r0)
        cantsplit(r0)
        for c, h, a, w_ in zip(r0.cells, headers, align, widths):
            cmargin(c, t=1.4, b=1.8, r=1.5)
            cborders(c, bottom=(8, P["RULE"]))
            valign(c, "bottom")
            run(cell_p(c, align=A[a]), h.upper(), font=F800, size=6.6,
                color=P["MUTED"], tr=track(6.6, 0.22))
        for ri, row in enumerate(rows):
            r = addrow(t, widths)
            cantsplit(r)
            for ci, (c, val) in enumerate(zip(r.cells, row)):
                cmargin(c, t=2.6, b=2.6, r=2.0)
                valign(c, "center" if ci in marks else "top")
                if ri < len(rows) - 1:
                    cborders(c, bottom=(2, P["PALE"]))
                p = cell_p(c, align=A[align[ci]], line=1.3)
                if ci in marks:
                    key = str(val).upper()
                    role = MARK_ROLE.get(key, "LBL")
                    col = self.LBL if role == "LBL" else P[role]
                    run(p, key, font=F800, size=6.8, color=col,
                        tr=track(6.8, 0.14))
                elif ci == 0 and first_bold:
                    run(p, _plain(val), font=F800, size=9.0, color=P["INK"])
                else:
                    rich(p, val, F400, 9.0, P["BODY"], bold_color=P["INK"])
        # keep the heading, the header row and the first data row together so
        # a band or label is never left alone at the foot of a page
        keeprow(r0)
        if rows:
            keeprow(t.rows[1])

    def item(self, question, label="POSITION GIVEN", answer="", min_h=9):
        """Numbered question with a labelled, shaded answer box."""
        P = self.P
        self.num += 1
        spacer(self.doc, 8, keep=True)
        t = table(self.doc, [GUTTER, BODY_W])
        r0 = t.rows[0]
        cantsplit(r0)
        valign(r0.cells[0], "top")
        cmargin(r0.cells[0], t=0.6)
        cmargin(r0.cells[1], l=0, t=0.6, b=1)
        run(cell_p(r0.cells[0], align=LEFT), "%02d" % self.num, font=F800,
            size=15, color=self.LBL, tr=track(15, -0.02))
        rich(cell_p(r0.cells[1], align=LEFT, line=1.3, keep=True, ind_r=4),
             question, F600, 10.6, P["INK"], F800)
        r1 = addrow(t, [GUTTER, BODY_W])
        cantsplit(r1)
        cmargin(r1.cells[1], l=0, t=2.8, b=1.4)
        run(cell_p(r1.cells[1], align=LEFT, keep=True), label.upper(),
            font=F800, size=6.8, color=P["MUTED"], tr=track(6.8, 0.22))
        keeprow(r0)
        keeprow(r1)
        r2 = addrow(t, [GUTTER, BODY_W])
        cantsplit(r2)
        rowheight(r2, min_h)
        box = r2.cells[1]
        shade(box, P["TINT"])
        cborders(box, top=(8, P["RULE"]), left=None, bottom=None, right=None)
        cmargin(box, l=5, r=5, t=2.5, b=2.5)
        rich(cell_p(box, align=LEFT, line=1.4), answer, F400, 9.4, P["BODY"],
             bold_color=P["INK"])

    def contents(self, label, rows, before=13):
        """Contents strip: (name, descriptor, count) rows under a small label."""
        P = self.P
        spacer(self.doc, before)
        run(para(self.doc), label.upper(), font=F800, size=8, color=self.LBL,
            tr=track(8, 0.20))
        spacer(self.doc, 4)
        cols = [43, 92, W - 135]
        t = table(self.doc, cols)
        for k, (name, sub, count) in enumerate(rows):
            r = t.rows[0] if k == 0 else addrow(t, cols)
            rowheight(r, 4.6)
            cantsplit(r)
            for c in r.cells:
                valign(c, "center")
                cmargin(c, t=2.2, b=2.2)
                if k < len(rows) - 1:
                    cborders(c, bottom=(2, P["TINT"]))
            cmargin(r.cells[2], r=1, t=2.2, b=2.2)
            run(cell_p(r.cells[0], align=LEFT), name, font=F800, size=10,
                color=P["INK"])
            run(cell_p(r.cells[1], align=LEFT), sub, font=F400, size=9,
                color=P["BODY"])
            run(cell_p(r.cells[2], align=RIGHT), count, font=F600, size=8.5,
                color=P["MUTED"])

    def notes_box(self, label, sub, height=27):
        """Free-text box (grows as the reader types): label, hint, shaded box."""
        P = self.P
        spacer(self.doc, 13, keep=True)
        t = table(self.doc, [GUTTER, BODY_W])
        r0 = t.rows[0]
        cantsplit(r0)
        cmargin(r0.cells[1], l=0, b=1.4)
        run(cell_p(r0.cells[1], align=LEFT), label.upper(), font=F800, size=8,
            color=self.LBL, tr=track(8, 0.20))
        run(cell_p(r0.cells[1], align=LEFT, before=2.5, line=1.4), sub,
            font=F400, size=9.2, color=P["BODY"])
        keeprow(r0)
        r1 = addrow(t, [GUTTER, BODY_W])
        cantsplit(r1)
        rowheight(r1, height)
        box = r1.cells[1]
        shade(box, P["TINT"])
        cborders(box, top=(8, P["RULE"]), left=None, bottom=None, right=None)
        cmargin(box, l=5, r=5, t=2.5, b=2.5)
        cell_p(box)

    def roadmap(self, phases, done=1, gap=2.0):
        """Row of phases: (title, sub). The first `done` are filled dark."""
        P, n = self.P, len(phases)
        w = (ROW_W - gap * (n - 1)) / n
        parts = []
        for i, (title, sub) in enumerate(phases):
            dark = i < done
            body = xpara([xrun("PHASE %d" % (i + 1), F800, 6.2,
                               PALE_OR(P, dark, "PALE", "META"), track(6.2, 0.2))],
                         after=2, jc="center")
            body += xpara([xrun(title, F800, 10, P["PAPER"] if dark else P["INK"],
                                track(10, -0.01))], after=1.5, jc="center",
                          line=1.1)
            body += xpara([xrun(sub, F400, 7, P["PALE"] if dark else P["BODY"])],
                          jc="center", line=1.15)
            if i:
                parts.append(shape(gap, 27.0, None, "<w:p/>", 0))
            parts.append(shape(w, 27.0, P["PANEL"] if dark else P["TINT"], body,
                               RADIUS, pad=(1.5, 2.6, 1.5, 2.6)))
        emit(self.doc, parts, before=12)

    def sign_row(self, cols, gap=3.0):
        """Signature boxes. cols = [(label, name, org)], name/org may be ''."""
        P, n = self.P, len(cols)
        w = (W - gap * (n - 1)) / n
        widths = []
        for i in range(n):
            widths += ([gap] if i else []) + [w]
        spacer(self.doc, 8)
        t = table(self.doc, widths)
        r = t.rows[0]
        rowheight(r, 27)
        cantsplit(r)
        for i, (label, name, org) in enumerate(cols):
            c = r.cells[i * 2]
            shade(c, P["TINT"])
            cborders(c, top=(8, self.LBL))
            cmargin(c, l=4, r=4, t=2.6, b=2.4)
            run(cell_p(c, align=LEFT), label.upper(), font=F800, size=6.6,
                color=P["MUTED"], tr=track(6.6, 0.22))
            line = cell_p(c, before=13, align=LEFT)
            pborder(line, "bottom", P["PANEL"], 6, 1)
            if name:
                run(cell_p(c, before=2, align=LEFT), name, font=F800, size=8.6,
                    color=P["INK"])
            else:
                cell_p(c, before=2, align=LEFT).add_run("").font.size = Pt(8.6)
            run(cell_p(c, align=LEFT), org, font=F400, size=7.6,
                color=P["MUTED"])
            run(cell_p(c, before=1.5, align=LEFT), "Date: ____________",
                font=F400, size=7.6, color=P["MUTED"])

    def contact_panel(self, label, name, line2=(), line3=(), prefix=""):
        """Dark contact panel. line2/line3 are lists of text parts joined
        with a pale middle dot. 38.5 mm tall so the last line never clips."""
        P = self.P
        spacer(self.doc, 14)

        def joined(parts, fonts):
            out = []
            for i, txt in enumerate(parts):
                if i:
                    out.append(xrun("   ·   ", F400, 8.8, P["PALE"]))
                out.append(xrun(txt, *fonts[min(i, len(fonts) - 1)]))
            return out
        body = xpara([xrun(label.upper(), F800, 6.8, P["PALE"], track(6.8, 0.22))],
                     after=4.5)
        body += xpara(([xrun(prefix, F500, 11, P["PALE"])] if prefix else []) +
                      [xrun(name, F800, 15.5, P["PAPER"], track(15.5, -0.02))],
                      after=2)
        if line2:
            body += xpara(joined(line2, [(F600, 8.8, P["TINT"]),
                                         (F500, 8.8, P["PALE"])]), after=3.5)
        if line3:
            body += xpara(joined(line3, [(F600, 8.8, P["TINT"]),
                                         (F400, 8.8, P["PALE"])]))
        emit(self.doc, [shape(W, 38.5, P["PANEL"], body, RADIUS,
                              pad=(7, 6.5, 7, 7))])

    def closing(self):
        out = os.path.join(self.tmp, "closing.png")
        self.b.closing(out)
        spacer(self.doc, 6)
        picture(para(self.doc, align=LEFT), out, width_mm=W)

    def page_break(self):
        para(self.doc).add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------- header / footer ------
    def _running(self):
        P, sec, b = self.P, self.sec, self.b
        # Header/Footer styles carry their own tab stops, which beat a
        # paragraph-level one, so both are laid out as borderless tables.
        sec.different_first_page_header_footer = True
        drop_empty_first(sec.header)
        ht = table(sec.header, [7.5, 47, W - 54.5])
        for c in ht.rows[0].cells:
            cborders(c, bottom=(4, P["PALE"]))
            cmargin(c, b=1.6)
            valign(c, "bottom")
        picture(cell_p(ht.rows[0].cells[0], align=LEFT), b.header_logo,
                height_mm=4.6)
        run(cell_p(ht.rows[0].cells[1], align=LEFT), b.wordmark, font=F800,
            size=6.8, color=P["INK"], tr=track(6.8, 0.14))
        p = cell_p(ht.rows[0].cells[2], align=RIGHT)
        run(p, self.kind, font=F600, size=7.2, color=self.LBL)
        run(p, "  ", font=F400, size=7.2, color=P["PALE"])
        run(p, self.title, font=F500, size=7.2, color=P["MUTED"])

        for foot in (sec.footer, sec.first_page_footer):
            drop_empty_first(foot)
            ft = table(foot, [W - 46, 46])
            for c in ft.rows[0].cells:
                cborders(c, top=(4, P["RULE"]))
                cmargin(c, t=1.8)
            run(cell_p(ft.rows[0].cells[0], align=LEFT), b.footer_tagline,
                font=F500, size=7, color=P["MUTED"])
            p = cell_p(ft.rows[0].cells[1], align=RIGHT)
            run(p, self.date, font=F600, size=6.6, color=P["MUTED"],
                tr=track(6.6, 0.12))
            run(p, "  ·  ", font=F400, size=7, color=P["PALE"])
            fld = el("w:fldSimple", instr=" PAGE ")
            sub = p.add_run("1")
            sub.font.size = Pt(7.5)
            sub.font.color.rgb = RGBColor.from_string(self.LBL)
            sub._element.get_or_add_rPr().insert(
                0, el("w:rFonts", ascii=F800, hAnsi=F800, cs=F800))
            fld.append(sub._element)
            p._p.append(fld)

    def save(self, out):
        self.doc.save(out)
        if self.embed:
            embed_fonts(out, FONT_DIR)
        return out


def PALE_OR(P, dark, a, b):
    return P[a] if dark else P[b]
